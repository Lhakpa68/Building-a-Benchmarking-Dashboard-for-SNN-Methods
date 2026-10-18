"""Renumber citations in order of first appearance and relink them.

python3 renumber.py in.docx out.docx

- Citation = a w:hyperlink whose visible text is "[n]".
- New numbers follow the order in which references are first cited in the document body.
- The reference list is reordered to match, its leading "[n]" labels rewritten, and every citation's
  hyperlink anchor points at a bookmark on the reference paragraph it cites (created if missing).
- Within a run of adjacent citations separated only by commas/spaces, citations are sorted ascending.
- Checks and reports: citations whose old link pointed at a different reference than their label,
  references never cited, plain-text "[n]" outside hyperlinks, and a final self-check.
"""
import sys, re, copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC, DST = sys.argv[1], sys.argv[2]
doc = Document(SRC)
body = doc.element.body
P = doc.paragraphs
LAB = re.compile(r'^\[(\d+)\]$')

def text_of(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))

# ---- reference list -------------------------------------------------------
ref_head = next(i for i, p in enumerate(P) if p.style.name.startswith('Heading 1') and p.text.strip() == 'REFERENCES')
ref_end = next(i for i, p in enumerate(P) if i > ref_head and p.style.name.startswith('Heading 1') and p.text.strip())
refs = {}          # old number -> paragraph element
ref_order = []
for i in range(ref_head + 1, ref_end):
    m = re.match(r'^\s*\[(\d+)\]', P[i].text)
    if m:
        n = int(m.group(1))
        assert n not in refs, f'duplicate reference number [{n}]'
        refs[n] = P[i]._p
        ref_order.append(n)
ref_elems = set(refs.values())
print(f'reference list: {len(refs)} entries, numbered {min(refs)}..{max(refs)}')

def bookmarks_in(p):
    return [b.get(qn('w:name')) for b in p.iter(qn('w:bookmarkStart'))]

bm_to_ref = {}
for n, p in refs.items():
    for b in bookmarks_in(p):
        bm_to_ref[b] = n

# ---- citations in document order (outside the reference list) --------------
cites = []
for h in body.iter(qn('w:hyperlink')):
    m = LAB.match(text_of(h).strip())
    if not m:
        continue
    anc = h.getparent()
    while anc is not None and anc.tag != qn('w:p'):
        anc = anc.getparent()
    if anc in ref_elems:
        continue
    cites.append((h, int(m.group(1))))
print(f'citations found: {len(cites)}')

report = []
for h, n in cites:
    a = h.get(qn('w:anchor'))
    if a is not None and a in bm_to_ref and bm_to_ref[a] != n:
        report.append(f'label [{n}] was linked to reference [{bm_to_ref[a]}] (fixed: now links to the reference its label names)')
    if n not in refs:
        report.append(f'label [{n}] has no reference entry')

# plain-text labels outside hyperlinks (not in the reference list)
for p in P[:ref_head] + P[ref_end:]:
    for r in p._p.findall(qn('w:r')):
        t = text_of(r)
        for m in re.finditer(r'\[(\d{1,3})\]', t):
            report.append(f'plain-text citation {m.group(0)} not hyperlinked: ...{p.text[max(0, p.text.find(m.group(0))-40):p.text.find(m.group(0))+10]}...')

# ---- new numbering ---------------------------------------------------------
new_of = {}
for h, n in cites:
    if n in refs and n not in new_of:
        new_of[n] = len(new_of) + 1
uncited = [n for n in ref_order if n not in new_of]
for n in uncited:
    new_of[n] = len(new_of) + 1
    report.append(f'reference [{n}] is never cited in the text (placed at the end as [{new_of[n]}])')

# bookmark per reference
bm_for = {}
for n, p in refs.items():
    names = [b for b in bookmarks_in(p) if not b.startswith('_GoBack')]
    if names:
        bm_for[n] = names[0]
    else:
        name = f'_RefCite{n:03d}'
        bid = str(900000 + n)
        s = OxmlElement('w:bookmarkStart'); s.set(qn('w:id'), bid); s.set(qn('w:name'), name)
        e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), bid)
        ppr = p.find(qn('w:pPr'))
        idx = list(p).index(ppr) + 1 if ppr is not None else 0
        p.insert(idx, s); p.insert(idx + 1, e)
        bm_for[n] = name

# relabel + relink citations
for h, n in cites:
    if n not in refs:
        continue
    ts = list(h.iter(qn('w:t')))
    full = ''.join(t.text or '' for t in ts)
    newlab = full.replace(f'[{n}]', f'[{new_of[n]}]', 1)
    ts[0].text = newlab
    for t in ts[1:]:
        t.text = ''
    h.set(qn('w:anchor'), bm_for[n])
    rid = h.get(qn('r:id'))
    if rid is not None:
        del h.attrib[qn('r:id')]
    h.set('_new', str(new_of[n]))

# sort adjacent groups ascending
SEP = re.compile(r'^[\s,;]*$')
groups_sorted = 0
for p in body.iter(qn('w:p')):
    kids = list(p)
    i = 0
    while i < len(kids):
        if kids[i].tag == qn('w:hyperlink') and kids[i].get('_new'):
            grp = [i]; j = i + 1
            while j < len(kids):
                k = kids[j]
                if k.tag == qn('w:hyperlink') and k.get('_new'):
                    grp.append(j); j += 1
                elif k.tag in (qn('w:r'), qn('w:proofErr'), qn('w:bookmarkStart'), qn('w:bookmarkEnd')) and SEP.match(text_of(k)):
                    j += 1
                else:
                    break
            if len(grp) > 1:
                els = [kids[g] for g in grp]
                srt = sorted(els, key=lambda e: int(e.get('_new')))
                if srt != els:
                    groups_sorted += 1
                    holders = []
                    for g in grp:
                        ph = OxmlElement('w:proofErr'); ph.set('_ph', '1')
                        p.replace(kids[g], ph); holders.append(ph)
                    for ph, e in zip(holders, srt):
                        p.replace(ph, e)
                    kids = list(p)
            i = j
        else:
            i += 1
for h in body.iter(qn('w:hyperlink')):
    if h.get('_new'):
        del h.attrib['_new']

# reorder the reference list and relabel
anchor = P[ref_head]._p
parent = anchor.getparent()
first_pos = min(list(parent).index(e) for e in refs.values())
for e in refs.values():
    parent.remove(e)
for k, n in enumerate(sorted(refs, key=lambda n: new_of[n])):
    e = refs[n]
    # the label is often split across runs ("[1" + "]  F. Bacho..."), so work on
    # the concatenated text of the leading runs
    ts = list(e.iter(qn('w:t')))
    joined = ''.join(t.text or '' for t in ts)
    m = re.search(r'\[\d+\]', joined)
    done = False
    if m:
        newjoined = joined[:m.start()] + f'[{new_of[n]}]' + joined[m.end():]
        # rewrite only as many leading runs as the label spans
        acc, idx = 0, []
        for j, t in enumerate(ts):
            ln = len(t.text or '')
            if acc < m.end() and acc + ln > m.start():
                idx.append(j)
            acc += ln
        first, last = idx[0], idx[-1]
        prefix = ''.join(ts[j].text or '' for j in range(first))
        tail_start = sum(len(ts[j].text or '') for j in range(last + 1))
        ts[first].text = newjoined[len(prefix):len(newjoined) - (len(joined) - tail_start)]
        for j in range(first + 1, last + 1):
            ts[j].text = ''
        done = True
    if not done:
        report.append(f'could not relabel reference [{n}]')
    parent.insert(first_pos + k, e)

doc.save(DST)

# ---- self-check ------------------------------------------------------------
d2 = Document(DST)
b2 = d2.element.body
bm = {}
for p in b2.iter(qn('w:p')):
    t = ''.join(x.text or '' for x in p.iter(qn('w:t')))
    m = re.match(r'^\s*\[(\d+)\]', t)
    if m:
        for b in p.iter(qn('w:bookmarkStart')):
            bm[b.get(qn('w:name'))] = int(m.group(1))
seq, bad, seen = [], 0, set()
for h in b2.iter(qn('w:hyperlink')):
    m = LAB.match(text_of(h).strip())
    if not m:
        continue
    n = int(m.group(1))
    a = h.get(qn('w:anchor'))
    if a not in bm or bm[a] != n:
        bad += 1
    if n not in seen:
        seen.add(n); seq.append(n)
first_order_ok = seq == sorted(seq) and seq == list(range(1, len(seq) + 1))
print(f'groups re-sorted: {groups_sorted}; citations linked to wrong entry after fix: {bad}; '
      f'first-appearance order ascending 1..{len(seq)}: {first_order_ok}')
print('\n'.join(report) if report else 'no issues found in the original citations')
