"""Rebuild the thesis with the rewritten paragraphs.

python3 apply.py src.docx final_text.json out.docx
- Rewritten paragraphs keep their paragraph properties, bookmarks and citation hyperlinks (deep copies of the
  originals, re-inserted at the ⟦Cn:..⟧ positions). Text runs take the paragraph's own character formatting and
  are highlighted yellow.
- Outside the rewritten paragraphs (headings, captions, tables, TOC), "framework" -> "simulator" is applied run by
  run, except in the reference list, the title/cover pages and proper names; every change is logged and
  highlighted.
"""
import sys, json, re, copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_COLOR_INDEX

SRC, TXT, DST = sys.argv[1:4]
doc = Document(SRC)
P = doc.paragraphs
new = {int(k): v for k, v in json.load(open(TXT)).items()}
CIT = re.compile(r'⟦C(\d+):([^⟧]*)⟧')


def yellow(rpr):
    h = rpr.find(qn('w:highlight'))
    if h is None:
        h = OxmlElement('w:highlight'); rpr.append(h)
    h.set(qn('w:val'), 'yellow')


def make_run(text, rpr_tpl):
    r = OxmlElement('w:r')
    rpr = copy.deepcopy(rpr_tpl) if rpr_tpl is not None else OxmlElement('w:rPr')
    yellow(rpr)
    r.append(rpr)
    t = OxmlElement('w:t'); t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    return r


problems = []
for i, text in sorted(new.items()):
    p = P[i]._p
    kids = list(p)
    ppr = p.find(qn('w:pPr'))
    marks = [c for c in kids if c.tag in (qn('w:bookmarkStart'), qn('w:bookmarkEnd'))]
    links = [c for c in kids if c.tag == qn('w:hyperlink')]
    tpl = None
    for c in kids:
        if c.tag == qn('w:r') and ''.join(t.text or '' for t in c.iter(qn('w:t'))).strip():
            tpl = c.find(qn('w:rPr')); break
    # check tokens against the original hyperlinks
    toks = CIT.findall(text)
    for n, shown in toks:
        n = int(n)
        if n < 1 or n > len(links):
            problems.append(f'[{i}] token C{n} has no hyperlink'); continue
        orig = ''.join(t.text or '' for t in links[n - 1].iter(qn('w:t')))
        if orig != shown:
            problems.append(f'[{i}] token C{n} text {shown!r} != hyperlink {orig!r}')
    if sorted(int(n) for n, _ in toks) != list(range(1, len(links) + 1)):
        problems.append(f'[{i}] citation set differs from original hyperlinks')
    for c in kids:
        if c is not ppr:
            p.remove(c)
    for m in marks:
        p.append(m)
    pos = 0
    for m in CIT.finditer(text):
        if m.start() > pos:
            p.append(make_run(text[pos:m.start()], tpl))
        link = copy.deepcopy(links[int(m.group(1)) - 1])
        for r in link.iter(qn('w:r')):
            rpr = r.find(qn('w:rPr'))
            if rpr is None:
                rpr = OxmlElement('w:rPr'); r.insert(0, rpr)
            yellow(rpr)
        p.append(link)
        pos = m.end()
    if pos < len(text):
        p.append(make_run(text[pos:], tpl))

# ---- terminology outside the rewritten paragraphs ----
ref_start = next(i for i, p in enumerate(P) if p.style.name.startswith('Heading 1') and p.text.strip() == 'REFERENCES')
ref_end = next(i for i, p in enumerate(P) if i > ref_start and p.style.name.startswith('Heading 1') and p.text.strip())
SKIP = set(range(0, 67)) | set(range(ref_start, ref_end)) | set(new)
RULES = [(re.compile(r'The SNN Software Framework Landscape'), 'The SNN Simulator Landscape'),
         (re.compile(r'SNN Software Framework Landscape'), 'SNN Simulator Landscape'),
         (re.compile(r'(?<!Engineering )\bFrameworks\b'), 'Simulators'),
         (re.compile(r'(?<!Engineering )\bframeworks\b'), 'simulators'),
         (re.compile(r'(?<!Engineering )\bFramework\b'), 'Simulator'),
         (re.compile(r'(?<!Engineering )\bframework\b'), 'simulator'),
         (re.compile(r'(?<!Engineering )\bFRAMEWORKS\b'), 'SIMULATORS')]
log = []


def fix_runs(par, where):
    for r in par.runs:
        t0 = r.text
        t = t0
        for rx, rep in RULES:
            t = rx.sub(rep, t)
        if t != t0:
            r.text = t
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            log.append(f'{where}: {t0!r} -> {t!r}')
    joined = par.text
    if re.search(r'(?<!Engineering )\b[Ff]rameworks?\b', joined):
        log.append(f'{where}: UNCHANGED (spans runs or exception): {joined[:120]!r}')


def span_replace(par, old, newtxt, where):
    """Replace text that is split across several adjacent runs."""
    from docx.text.run import Run
    runs = [Run(r, par) for r in par._p.iter(qn('w:r'))]   # includes runs inside hyperlinks
    joined = ''.join(r.text for r in runs)
    pos = joined.find(old)
    if pos < 0:
        return False
    starts, acc = [], 0
    for r in runs:
        starts.append(acc); acc += len(r.text)
    i0 = max(i for i, s in enumerate(starts) if s <= pos)
    end = pos + len(old)
    i1 = max(i for i, s in enumerate(starts) if s < end)
    merged = ''.join(r.text for r in runs[i0:i1 + 1])
    local = pos - starts[i0]
    runs[i0].text = merged[:local] + newtxt + merged[local + len(old):]
    runs[i0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    for r in runs[i0 + 1:i1 + 1]:
        r._r.getparent().remove(r._r)
    log.append(f'{where}: (across runs) {old!r} -> {newtxt!r}')
    return True


for i, par in enumerate(P):
    if i in SKIP:
        continue
    fix_runs(par, f'para {i}')
# the keywords line keeps its bold lead-in, so it is handled by name
span_replace(P[86], 'framework comparison', 'simulator comparison', 'para 86')
for ti, tbl in enumerate(doc.tables):
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            for par in cell.paragraphs:
                fix_runs(par, f'table {ti} r{ri} c{ci}')

TOC_FIX = [('The SNN Software Framework Landscape', 'The SNN Simulator Landscape'),
           ('Comparison against the fastest framework configuration', 'Comparison against the fastest simulator configuration'),
           ('One Row Per Framework/Backend Combination', 'One Row Per Simulator/Backend Combination'),
           ('the live framework selector', 'the live simulator selector')]
for ti, tbl in enumerate(doc.tables):
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            for par in cell.paragraphs:
                for old, newtxt in TOC_FIX:
                    span_replace(par, old, newtxt, f'table {ti} r{ri} c{ci}')

doc.save(DST)
open(DST + '.terminology.log', 'w').write('\n'.join(log))
print('rewritten paragraphs:', len(new))
print('terminology changes:', sum(1 for l in log if '->' in l), '| unchanged mentions:', sum(1 for l in log if 'UNCHANGED' in l))
print('\n'.join(problems) if problems else 'citations OK')
